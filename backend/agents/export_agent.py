import io
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, Paragraph, Spacer
from docx import Document


def generate_excel(
    transactions: list[dict[str, Any]],
    budgets: list[dict[str, Any]],
    summary: str = "",
) -> bytes:
    wb = Workbook()

    ws_tx = wb.active
    ws_tx.title = "Transactions"
    headers = ["Date", "Description", "Category", "Amount", "Currency", "Merchant"]
    ws_tx.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws_tx.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for tx in transactions:
        ws_tx.append([
            tx.get("date", ""),
            tx.get("description", ""),
            tx.get("category", ""),
            float(tx.get("amount", 0)),
            tx.get("currency", "USD"),
            tx.get("merchant", ""),
        ])

    ws_budget = wb.create_sheet("Budgets")
    budget_headers = ["Category", "Planned", "Actual", "Remaining"]
    ws_budget.append(budget_headers)
    for col in range(1, len(budget_headers) + 1):
        cell = ws_budget.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")

    for b in budgets:
        planned = float(b.get("planned_amount", 0))
        actual = float(b.get("actual_amount", 0))
        ws_budget.append([
            b.get("category", ""),
            planned,
            actual,
            planned - actual,
        ])

    if summary:
        ws_summary = wb.create_sheet("AI Summary")
        ws_summary["A1"] = "AI-Generated Financial Summary"
        ws_summary["A1"].font = Font(bold=True, size=14)
        ws_summary["A2"] = summary

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_pdf(
    transactions: list[dict[str, Any]],
    budgets: list[dict[str, Any]],
    summary: str = "",
) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter)
    styles = getSampleStyleSheet()
    elements: list[Any] = []

    elements.append(Paragraph("Spentra Financial Report", styles["Title"]))
    elements.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d')}", styles["Normal"]))
    elements.append(Spacer(1, 20))

    if summary:
        elements.append(Paragraph("AI Summary", styles["Heading2"]))
        elements.append(Paragraph(summary, styles["Normal"]))
        elements.append(Spacer(1, 20))

    elements.append(Paragraph("Transactions", styles["Heading2"]))
    tx_data = [["Date", "Description", "Category", "Amount", "Currency"]]
    for tx in transactions:
        tx_data.append([
            str(tx.get("date", "")),
            str(tx.get("description", "")),
            str(tx.get("category", "")),
            str(tx.get("amount", "")),
            str(tx.get("currency", "USD")),
        ])
    elements.append(Table(tx_data))
    elements.append(Spacer(1, 20))

    elements.append(Paragraph("Budgets", styles["Heading2"]))
    budget_data = [["Category", "Planned", "Actual", "Remaining"]]
    for b in budgets:
        planned = float(b.get("planned_amount", 0))
        actual = float(b.get("actual_amount", 0))
        budget_data.append([
            str(b.get("category", "")),
            str(planned),
            str(actual),
            str(planned - actual),
        ])
    elements.append(Table(budget_data))

    doc.build(elements)
    return buf.getvalue()


def generate_word(
    transactions: list[dict[str, Any]],
    budgets: list[dict[str, Any]],
    summary: str = "",
) -> bytes:
    doc = Document()
    doc.add_heading("Spentra Financial Report", 0)
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")

    if summary:
        doc.add_heading("AI Summary", level=1)
        doc.add_paragraph(summary)

    doc.add_heading("Transactions", level=1)
    tx_table = doc.add_table(rows=1, cols=5)
    tx_table.style = "Table Grid"
    hdr = tx_table.rows[0].cells
    for i, h in enumerate(["Date", "Description", "Category", "Amount", "Currency"]):
        hdr[i].text = h
    for tx in transactions:
        row = tx_table.add_row().cells
        row[0].text = str(tx.get("date", ""))
        row[1].text = str(tx.get("description", ""))
        row[2].text = str(tx.get("category", ""))
        row[3].text = str(tx.get("amount", ""))
        row[4].text = str(tx.get("currency", "USD"))

    doc.add_heading("Budgets", level=1)
    b_table = doc.add_table(rows=1, cols=4)
    b_table.style = "Table Grid"
    hdr = b_table.rows[0].cells
    for i, h in enumerate(["Category", "Planned", "Actual", "Remaining"]):
        hdr[i].text = h
    for b in budgets:
        planned = float(b.get("planned_amount", 0))
        actual = float(b.get("actual_amount", 0))
        row = b_table.add_row().cells
        row[0].text = str(b.get("category", ""))
        row[1].text = str(planned)
        row[2].text = str(actual)
        row[3].text = str(planned - actual)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
